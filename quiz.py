import sqlite3
import pandas
import openpyxl


def initialize_quiz_database():  # first time start up, RUN THIS IF YOU HAVE NOT INITIALIZED THE QUIZ DATABASES ON YOUR
    # SYSTEM
    conn = sqlite3.connect('quiz.db')
    cursor = conn.cursor()
    try:
        cursor.execute("""
                CREATE TABLE leaderboard (
                    user_id INTEGER PRIMARY KEY,
                    username TEXT,
                    score INTEGER
                );
            """)
    except sqlite3.OperationalError as e:
        print("Error, leaderboard table already exists!")
    except Exception as e:
        print(f"Unknown error: {e}")
    try:
        cursor.execute("""
                CREATE TABLE masterQuiz (
                    quiz_id INTEGER PRIMARY KEY,
                    quiz_name TEXT,
                    quiz_display_name TEXT,
                    quiz_description TEXT
                );
            """)
    except sqlite3.OperationalError as e:
        print("Error, master quiz table already exists!")
    except Exception as e:
        print(f"Unknown error: {e}")
    conn.commit()
    conn.close()


class QuizEdit:
    def __init__(self):
        self.conn = sqlite3.connect('quiz.db')
        self.cursor = self.conn.cursor()

    def create_table(self, name, column_definitions, displayname, description):
        sql = f"CREATE TABLE {name} ({column_definitions});"
        try:
            self.cursor.execute(sql)
            self.conn.commit()
            print(f"Table '{name}' has been created!")
        except sqlite3.Error as e:
            print(f"Error creating table '{name}': {e}")
        try:
            self.cursor.execute('SELECT COUNT(*) FROM masterQuiz')
            row_count = self.cursor.fetchone()[0]
            sql = f"INSERT INTO masterQuiz (quiz_id, quiz_name, quiz_display_name, quiz_description) VALUES ({row_count}, '{name}', '{displayname}', '{description}')"
            self.cursor.execute(sql)
            self.conn.commit()
        except Exception as e:
            print(f"Error creating table key: {e}")

    def add_question(self, table_name, data): #add question singularly
        placeholders = ", ".join(['?'] * len(data))
        try:
            sql = f"INSERT INTO {table_name} VALUES ({placeholders})"
            self.cursor.execute(sql, tuple(data))
            self.conn.commit()
        except sqlite3.Error as e:
            print(f"Error inserting data into table '{table_name}: {e}'")
        except Exception as e:
            print(f"Unknown error: {e}")

    def add_question_excel(self, name, excel_file_name): #add multiple questions at once via excel file
        if '.xlsx' in excel_file_name:
            excel_file_name = excel_file_name.removesuffix('.xlsx')
        try:
            df = pandas.read_excel(f'excel/{excel_file_name}.xlsx', header=None)
        except Exception as e:
            print(f"Error: {e}")
            return

        first_row = df.iloc[0]

        header_words = ['question', 'answer']

        if any(word.lower() in ' '.join(first_row.str.lower()) for word in header_words): # checks if the header_words are contained in the first row
            df = df[1:] #excludes the first row

        try:
            for row in df.itertuples(index=False):
                placeholders = ', '.join(['?'] * len(row))
                sql = f"INSERT INTO {name} VALUES ({placeholders})"
                self.cursor.execute(sql, row)
            self.conn.commit()
            print(f"Data for {name} from {excel_file_name} inserted successfully!")
        except sqlite3.Error as e:
            print(f"Error inserting data into table '{name}: {e}'")
        except Exception as e:
            print(f"Unknown error: {e}")

    def close(self):
        self.conn.close()
initialize_quiz_database()
db = QuizEdit()
db.create_table("quiz_uscapitals", "question TEXT, answer TEXT", "US Capitals",
                "This quiz set contains all 50 United States Captials and States!")
db.add_question_excel("quiz_uscapitals", "quiz_uscapitals")
#db.add_question("quiz_test", ["What is the color of the sky?", "blue"]) format for adding questions
db.close()
